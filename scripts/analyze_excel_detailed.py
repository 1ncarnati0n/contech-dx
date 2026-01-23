#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Excel 파일 시트별 상세 구조 분석 스크립트
"""
import sys
import os
import io
import json

# Windows에서 UTF-8 출력을 위한 설정
if sys.platform == 'win32':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

try:
    import openpyxl
except ImportError:
    print("Error: openpyxl 라이브러리가 필요합니다.")
    print("설치 방법: pip install openpyxl")
    sys.exit(1)

def safe_str(value, max_len=50):
    """값을 안전하게 문자열로 변환"""
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return str(value)
    s = str(value)
    if len(s) > max_len:
        return s[:max_len] + "..."
    return s

def analyze_sheet(ws, sheet_name):
    """시트 상세 분석"""
    result = {
        "sheet_name": sheet_name,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "dimensions": ws.dimensions,
        "headers": [],
        "data_rows": [],
        "formulas": []
    }
    
    # 헤더 행 찾기 (비어있지 않은 첫 번째 행)
    header_row = None
    for row_idx in range(1, min(11, ws.max_row + 1)):
        row_values = [safe_str(cell.value) for cell in ws[row_idx]]
        if any(v.strip() for v in row_values):
            if header_row is None:
                header_row = row_idx
            result["headers"].append({
                "row": row_idx,
                "values": row_values
            })
    
    # 데이터 행 (헤더 이후)
    start_data_row = (header_row or 1) + 1
    for row_idx in range(start_data_row, min(start_data_row + 20, ws.max_row + 1)):
        row_values = []
        row_formulas = []
        for col_idx in range(1, min(ws.max_column + 1, 15)):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell_address = openpyxl.utils.get_column_letter(col_idx) + str(row_idx)
            
            if cell.data_type == 'f':  # formula
                formula = f"{cell_address}={cell.value}"
                row_formulas.append(formula)
                row_values.append({
                    "address": cell_address,
                    "value": safe_str(cell.value),
                    "formula": cell.value
                })
            else:
                row_values.append({
                    "address": cell_address,
                    "value": safe_str(cell.value)
                })
        
        if any(v["value"].strip() for v in row_values):
            result["data_rows"].append({
                "row": row_idx,
                "values": row_values,
                "formulas": row_formulas
            })
    
    # 모든 수식 찾기
    for row in ws.iter_rows():
        for cell in row:
            if cell.data_type == 'f':
                result["formulas"].append({
                    "address": cell.coordinate,
                    "formula": cell.value
                })
    
    return result

def main():
    file_path = "docs/작업면작업면적.xlsx"
    
    if not os.path.exists(file_path):
        print(f"Error: 파일을 찾을 수 없습니다: {file_path}")
        sys.exit(1)
    
    wb = openpyxl.load_workbook(file_path, data_only=False)
    
    all_results = {
        "file_name": os.path.basename(file_path),
        "sheet_count": len(wb.sheetnames),
        "sheet_names": wb.sheetnames,
        "sheets": []
    }
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        result = analyze_sheet(ws, sheet_name)
        all_results["sheets"].append(result)
    
    # JSON으로 출력
    output_file = "scripts/excel_analysis.json"
    with open(output_file, 'w', encoding='utf-8') as f:
        json.dump(all_results, f, ensure_ascii=False, indent=2)
    
    print(f"분석 완료. 결과가 {output_file}에 저장되었습니다.")
    
    # 간단한 요약 출력
    print(f"\n파일명: {all_results['file_name']}")
    print(f"시트 개수: {all_results['sheet_count']}")
    print(f"\n시트 목록:")
    for i, sheet_name in enumerate(all_results['sheet_names'], 1):
        sheet_data = all_results['sheets'][i-1]
        print(f"  {i}. {sheet_name} (행: {sheet_data['max_row']}, 열: {sheet_data['max_column']})")

if __name__ == "__main__":
    main()
