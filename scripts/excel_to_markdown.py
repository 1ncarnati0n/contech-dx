#!/usr/bin/env python3
"""
Excel 시트를 단일 마크다운 파일로 변환하는 스크립트 (개선 버전)
StandardProcess_Input_Data.xlsx → StandardProcess_Input_Data.md

개선 사항:
- 다중 헤더 행 처리 (병합 셀 forward-fill)
- 부동소수점 정밀도 오류 수정 (소수점 4자리 반올림)
- 빈 컬럼명 자동 생성
"""

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pathlib import Path
from typing import Optional

# 파일 경로 설정
EXCEL_FILE = Path(__file__).parent.parent / "docs" / "StandardProcess_Input_Data.xlsx"
OUTPUT_FILE = Path(__file__).parent.parent / "docs" / "StandardProcess_Input_Data.md"

# 숫자 반올림 자릿수
FLOAT_PRECISION = 4

# 시트명 → 한글 설명 매핑
SHEET_DESCRIPTIONS = {
    "지상 골조표준공정": "지상 골조 표준공정",
    "주동지하 표준공정": "주동 지하층 표준공정",
    "주차장 표준공정": "주차장 표준공정",
    "주차장  표준공정": "주차장 표준공정",
    "층고6.5m이상구간": "층고 6.5m 이상 구간",
    "층고": "층고 정보",
    "작업조구성": "작업조 구성 정보",
    "공종별 인원 상,하한 설정": "공종별 인원 한계 설정",
    "부분별 보정계수설정": "부분별 보정계수",
    "거푸집 계열 작업면, 작업면적 추정 산식": "거푸집 작업면적 산식",
    "철근 작업면,작업면적(Zone,Front) 추정 산식": "철근 작업면적 산식",
    "인당생산성": "인당 생산성 데이터",
    "펌프카 부위별 타설량": "펌프카 타설량",
}

# 시트별 헤더 행 수 설정 (기본값: 1)
SHEET_HEADER_ROWS = {
    "지상 골조표준공정": 2,
    "주동지하 표준공정": 2,
    "주차장  표준공정": 2,
    "층고6.5m이상구간": 2,
    "인당생산성": 1,
    "펌프카 부위별 타설량": 0,  # 헤더 없음
}

# 시트별 헤더 교체 설정
# skip_first_row: True면 기존 헤더 행들을 스킵하고 header_row_index 행을 헤더로 사용
# header_row_index: 헤더로 사용할 행 (0-indexed)
# category_row_index: 헤더와 데이터 사이에 포함할 카테고리 행 (0-indexed, optional)
# remove_empty_cols: True면 빈 컬럼 제거
SHEET_HEADER_CONFIG = {
    "지상 골조표준공정": {
        "skip_first_row": True,      # Row 1 스킵 (빈 값)
        "header_row_index": 2,       # Row 3을 헤더로 사용 (0-indexed: 2)
        "category_row_index": 1,     # Row 2를 카테고리 행으로 포함 (참고자료, 작업면적기준 등)
        "remove_empty_cols": True,   # 빈 컬럼 제거
    },
    "주동지하 표준공정": {
        "skip_first_row": True,
        "header_row_index": 2,
        "category_row_index": 1,
        "remove_empty_cols": True,
    },
    "주차장  표준공정": {
        "skip_first_row": True,
        "header_row_index": 2,
        "category_row_index": 1,
        "remove_empty_cols": True,
    },
    "층고6.5m이상구간": {
        "skip_first_row": True,
        "header_row_index": 2,
        "category_row_index": 1,
        "remove_empty_cols": True,
    },
}


def round_float(value) -> str:
    """부동소수점 값을 안전하게 문자열로 변환"""
    if pd.isna(value):
        return ""

    if isinstance(value, float):
        # 정수로 표현 가능하면 정수로
        if value == int(value):
            return str(int(value))
        # 소수점 반올림
        rounded = round(value, FLOAT_PRECISION)
        # 불필요한 0 제거
        return f"{rounded:g}"

    return str(value)


def get_merged_cells_map(ws) -> dict:
    """병합 셀 정보를 딕셔너리로 반환 {(row, col): value}"""
    merged_map = {}

    for merged_range in ws.merged_cells.ranges:
        min_row, min_col = merged_range.min_row, merged_range.min_col
        max_row, max_col = merged_range.max_row, merged_range.max_col
        value = ws.cell(min_row, min_col).value

        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                merged_map[(row, col)] = value

    return merged_map


def remove_empty_columns(df: pd.DataFrame) -> pd.DataFrame:
    """빈 컬럼 제거 (헤더가 비어있거나 모든 데이터가 빈 컬럼)"""
    cols_to_keep = []

    # 인덱스 기반으로 처리 (중복 컬럼명 대응)
    for col_idx, col in enumerate(df.columns):
        col_data = df.iloc[:, col_idx]

        # 헤더가 빈 문자열이거나 "Col숫자" 패턴이면 제거 대상
        col_str = str(col).strip()
        if not col_str or (col_str.startswith("Col") and col_str[3:].isdigit()):
            # 데이터도 모두 비어있는지 확인
            if col_data.isna().all() or (col_data.astype(str).str.strip() == "").all():
                continue  # 제거

        # 모든 데이터가 비어있으면 제거
        if col_data.isna().all():
            continue
        if (col_data.astype(str).str.strip() == "").all():
            continue

        cols_to_keep.append(col_idx)

    return df.iloc[:, cols_to_keep]


def read_sheet_with_headers(wb, sheet_name: str, header_rows: int = 1) -> pd.DataFrame:
    """시트를 읽고 다중 헤더 처리"""
    ws = wb[sheet_name]
    merged_map = get_merged_cells_map(ws)

    # 시트별 특수 설정 확인
    header_config = SHEET_HEADER_CONFIG.get(sheet_name, {})
    skip_first_row = header_config.get("skip_first_row", False)
    header_row_index = header_config.get("header_row_index", 0)
    category_row_index = header_config.get("category_row_index", None)
    should_remove_empty_cols = header_config.get("remove_empty_cols", False)

    # 전체 데이터 읽기
    data = []
    max_col = ws.max_column
    max_row = ws.max_row

    for row_idx in range(1, max_row + 1):
        row_data = []
        for col_idx in range(1, max_col + 1):
            # 병합 셀이면 병합된 값 사용
            if (row_idx, col_idx) in merged_map:
                value = merged_map[(row_idx, col_idx)]
            else:
                value = ws.cell(row_idx, col_idx).value
            row_data.append(value)
        data.append(row_data)

    if not data:
        return pd.DataFrame()

    # ===== 특수 헤더 처리: Row 스킵 → 지정된 행을 헤더로 사용 =====
    if skip_first_row and len(data) > header_row_index:
        # 지정된 행을 헤더로 사용
        headers = data[header_row_index]

        # 카테고리 행 포함 여부에 따라 body_data 구성
        if category_row_index is not None and category_row_index < len(data):
            # 카테고리 행을 데이터의 첫 번째 행으로 포함
            category_row = data[category_row_index]
            body_data = [category_row] + data[header_row_index + 1:]
        else:
            body_data = data[header_row_index + 1:]

        # 헤더 정리 (None, 빈 문자열 처리)
        cleaned_headers = []
        unnamed_count = 0
        for i, h in enumerate(headers):
            if h is None or str(h).strip() == "":
                unnamed_count += 1
                cleaned_headers.append(f"Col{i+1}")
            else:
                cleaned_headers.append(str(h).strip())

        df = pd.DataFrame(body_data, columns=cleaned_headers)

        # 빈 컬럼 제거
        if should_remove_empty_cols:
            df = remove_empty_columns(df)

        return df

    # 헤더가 없는 경우
    if header_rows == 0:
        df = pd.DataFrame(data)
        df.columns = [f"Col{i+1}" for i in range(len(df.columns))]
        return df

    # 단일 헤더
    if header_rows == 1:
        headers = data[0]
        df = pd.DataFrame(data[1:], columns=headers)
    # 다중 헤더 (2행 이상)
    else:
        # 헤더 행들을 결합
        header_data = data[:header_rows]
        body_data = data[header_rows:]

        # 각 컬럼의 헤더를 병합 (빈 값은 위 행에서 forward-fill)
        combined_headers = []
        for col_idx in range(max_col):
            parts = []
            last_non_empty = ""
            for row_idx in range(header_rows):
                val = header_data[row_idx][col_idx]
                if val is not None and str(val).strip():
                    last_non_empty = str(val).strip()
                    parts.append(last_non_empty)
                elif last_non_empty and row_idx > 0:
                    # 빈 값이면 이전 행 값 유지 (forward-fill 효과)
                    pass
                else:
                    parts.append("")

            # 중복 제거 후 결합
            unique_parts = []
            for p in parts:
                if p and p not in unique_parts:
                    unique_parts.append(p)

            if unique_parts:
                combined_headers.append(" / ".join(unique_parts))
            else:
                combined_headers.append(f"Col{col_idx+1}")

        df = pd.DataFrame(body_data, columns=combined_headers)

    # 컬럼명 정리 (None, 빈 문자열 처리)
    new_columns = []
    unnamed_count = 0
    for i, col in enumerate(df.columns):
        if col is None or str(col).strip() == "" or str(col).startswith("Unnamed:"):
            unnamed_count += 1
            new_columns.append(f"항목{unnamed_count}")
        else:
            new_columns.append(str(col).strip())
    df.columns = new_columns

    # 빈 컬럼 제거 (설정된 경우)
    if should_remove_empty_cols:
        df = remove_empty_columns(df)

    return df


def dataframe_to_markdown_table(df: pd.DataFrame) -> str:
    """DataFrame을 마크다운 테이블로 변환 (부동소수점 정밀도 처리)"""
    if df.empty:
        return "_데이터 없음_"

    # 컬럼명 처리
    headers = [str(col).replace("|", "\\|") for col in df.columns]

    # 헤더 행
    header_row = "| " + " | ".join(headers) + " |"

    # 구분선
    separator = "| " + " | ".join(["---"] * len(headers)) + " |"

    # 데이터 행
    data_rows = []
    for _, row in df.iterrows():
        cells = []
        for cell in row:
            cell_str = round_float(cell)
            cell_str = cell_str.replace("|", "\\|").replace("\n", " ")
            cells.append(cell_str)
        data_rows.append("| " + " | ".join(cells) + " |")

    return "\n".join([header_row, separator] + data_rows)


def convert_sheet_to_section(sheet_num: int, sheet_name: str, df: pd.DataFrame) -> str:
    """시트를 마크다운 섹션으로 변환"""
    row_count = len(df)
    col_count = len(df.columns)
    description = SHEET_DESCRIPTIONS.get(sheet_name, sheet_name)

    section = f"""## {sheet_num}. {sheet_name}

> {description}

- **데이터**: {row_count}행 × {col_count}열

{dataframe_to_markdown_table(df)}

---
"""
    return section


def main():
    print(f"📂 Excel 파일 읽기: {EXCEL_FILE}")

    # openpyxl로 워크북 로드 (병합 셀 처리용)
    wb = load_workbook(EXCEL_FILE, data_only=True)
    available_sheets = wb.sheetnames

    print(f"📋 발견된 시트 ({len(available_sheets)}개):")
    for i, name in enumerate(available_sheets, 1):
        header_rows = SHEET_HEADER_ROWS.get(name, 1)
        print(f"   {i}. {name} (헤더: {header_rows}행)")

    print("\n" + "=" * 50)

    # 문서 헤더 생성
    md_parts = [
        "# Standard Process Input Data",
        "",
        "## 개요",
        "",
        "- **원본 파일**: `StandardProcess_Input_Data.xlsx`",
        f"- **시트 수**: {len(available_sheets)}개",
        "",
        "## 목차",
        "",
    ]

    # 목차 생성
    for i, sheet_name in enumerate(available_sheets, 1):
        description = SHEET_DESCRIPTIONS.get(sheet_name, sheet_name)
        md_parts.append(f"{i}. [{sheet_name}](#{sheet_name.replace(' ', '-').lower()}) - {description}")

    md_parts.append("")
    md_parts.append("---")
    md_parts.append("")

    # 각 시트를 섹션으로 변환
    for i, sheet_name in enumerate(available_sheets, 1):
        header_rows = SHEET_HEADER_ROWS.get(sheet_name, 1)
        df = read_sheet_with_headers(wb, sheet_name, header_rows)

        # 빈 행 제거 (모든 값이 빈 행)
        df = df.dropna(how='all')

        section = convert_sheet_to_section(i, sheet_name, df)
        md_parts.append(section)
        print(f"✅ {i}. {sheet_name} ({len(df)}행)")

    # 단일 파일로 저장
    full_content = "\n".join(md_parts)
    OUTPUT_FILE.write_text(full_content, encoding="utf-8")

    wb.close()

    print("\n" + "=" * 50)
    print(f"🎉 완료: 단일 파일 생성됨 (개선 버전)")
    print(f"📁 출력 파일: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
