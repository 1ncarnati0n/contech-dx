#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Excel 파일 시트별 구조 분석 스크립트
"""
import sys
import os
import io

# Windows에서 UTF-8 출력을 위한 설정
if sys.platform == 'win32':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

try:
    import openpyxl
    USE_OPENPYXL = True
except ImportError:
    try:
        import pandas as pd
        USE_PANDAS = True
    except ImportError:
        print("Error: openpyxl 또는 pandas 라이브러리가 필요합니다.")
        print("설치 방법: pip install openpyxl 또는 pip install pandas openpyxl")
        sys.exit(1)

def analyze_with_openpyxl(file_path):
    """openpyxl을 사용하여 Excel 파일 분석"""
    wb = openpyxl.load_workbook(file_path, data_only=False)
    
    print(f"\n파일명: {os.path.basename(file_path)}")
    print(f"시트 개수: {len(wb.sheetnames)}")
    print(f"시트 목록: {', '.join(wb.sheetnames)}")
    print("\n" + "="*80 + "\n")
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"시트명: {sheet_name}")
        print(f"최대 행: {ws.max_row}")
        print(f"최대 열: {ws.max_column}")
        print(f"사용된 범위: {ws.dimensions}")
        print("\n--- 헤더 구조 (상위 5행) ---")
        
        # 상위 5행 출력
        for row_idx in range(1, min(6, ws.max_row + 1)):
            row_data = []
            for col_idx in range(1, min(ws.max_column + 1, 15)):  # 최대 14열까지
                cell = ws.cell(row=row_idx, column=col_idx)
                cell_value = cell.value
                if cell_value is None:
                    cell_value = ""
                elif isinstance(cell_value, (int, float)):
                    cell_value = str(cell_value)
                else:
                    cell_value = str(cell_value)[:30]  # 30자로 제한
                
                # 셀 주소와 값
                cell_address = openpyxl.utils.get_column_letter(col_idx) + str(row_idx)
                row_data.append(f"{cell_address}:{cell_value}")
            
            if any(cell.value for cell in ws[row_idx] if cell.value):
                print(f"행 {row_idx}: {' | '.join(row_data)}")
        
        print("\n--- 데이터 샘플 (6-15행) ---")
        for row_idx in range(6, min(16, ws.max_row + 1)):
            row_data = []
            for col_idx in range(1, min(ws.max_column + 1, 15)):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell_value = cell.value
                if cell_value is None:
                    cell_value = ""
                elif isinstance(cell_value, (int, float)):
                    cell_value = str(cell_value)
                else:
                    cell_value = str(cell_value)[:30]
                
                cell_address = openpyxl.utils.get_column_letter(col_idx) + str(row_idx)
                row_data.append(f"{cell_address}:{cell_value}")
            
            if any(cell.value for cell in ws[row_idx] if cell.value):
                print(f"행 {row_idx}: {' | '.join(row_data)}")
        
        # 수식이 있는 셀 찾기
        formula_cells = []
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == 'f':  # formula
                    formula_cells.append(f"{cell.coordinate}={cell.value}")
        
        if formula_cells:
            print(f"\n--- 수식 셀 (처음 10개) ---")
            for formula in formula_cells[:10]:
                print(f"  {formula}")
            if len(formula_cells) > 10:
                print(f"  ... 외 {len(formula_cells) - 10}개")
        
        print("\n" + "="*80 + "\n")

def analyze_with_pandas(file_path):
    """pandas를 사용하여 Excel 파일 분석"""
    xl_file = pd.ExcelFile(file_path)
    
    print(f"\n파일명: {os.path.basename(file_path)}")
    print(f"시트 개수: {len(xl_file.sheet_names)}")
    print(f"시트 목록: {', '.join(xl_file.sheet_names)}")
    print("\n" + "="*80 + "\n")
    
    for sheet_name in xl_file.sheet_names:
        df = pd.read_excel(file_path, sheet_name=sheet_name, header=None)
        print(f"시트명: {sheet_name}")
        print(f"행 수: {len(df)}")
        print(f"열 수: {len(df.columns)}")
        print("\n--- 상위 10행 데이터 ---")
        print(df.head(10).to_string())
        print("\n" + "="*80 + "\n")

if __name__ == "__main__":
    file_path = "docs/작업면작업면적.xlsx"
    
    if not os.path.exists(file_path):
        print(f"Error: 파일을 찾을 수 없습니다: {file_path}")
        sys.exit(1)
    
    if USE_OPENPYXL:
        analyze_with_openpyxl(file_path)
    elif USE_PANDAS:
        analyze_with_pandas(file_path)
