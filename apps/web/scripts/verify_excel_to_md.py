#!/usr/bin/env python3
"""
Excel ↔ Markdown 데이터 정합성 검증 스크립트

검증 항목:
1. 시트별 행/열 개수 비교
2. 셀 값 샘플링 비교
3. 숫자값 정밀도 검증
4. 헤더 손실 감지

Usage:
    python scripts/verify_excel_to_md.py
"""

import pandas as pd
import re
import json
from pathlib import Path
from dataclasses import dataclass, field
from typing import Optional
import random
from openpyxl import load_workbook

# 파일 경로
BASE_DIR = Path(__file__).parent.parent
EXCEL_FILE = BASE_DIR / "docs" / "StandardProcess_Input_Data.xlsx"
MD_FILE = BASE_DIR / "docs" / "StandardProcess_Input_Data.md"
REPORT_FILE = BASE_DIR / "docs" / "verification_report.md"

# 검증 설정
NUMERIC_TOLERANCE = 0.001  # 숫자 비교 허용 오차
SAMPLE_ROWS_PER_SHEET = 5  # 시트당 샘플링할 행 수
FLOAT_PRECISION = 4  # 소수점 반올림 자릿수

# 시트별 헤더 행 수 설정 (excel_to_markdown.py와 동일)
SHEET_HEADER_ROWS = {
    "지상 골조표준공정": 2,
    "주동지하 표준공정": 2,
    "주차장  표준공정": 2,
    "층고6.5m이상구간": 2,
    "인당생산성": 1,
    "펌프카 부위별 타설량": 0,  # 헤더 없음
}

# 시트별 헤더 교체 설정 (excel_to_markdown.py와 동일)
# category_row_index: 헤더와 데이터 사이에 포함할 카테고리 행 (0-indexed, optional)
SHEET_HEADER_CONFIG = {
    "지상 골조표준공정": {
        "skip_first_row": True,
        "header_row_index": 2,       # Row 3을 헤더로 사용 (0-indexed)
        "category_row_index": 1,     # Row 2를 카테고리 행으로 포함
        "remove_empty_cols": True,
    },
    "주동지하 표준공정": {
        "skip_first_row": True,
        "header_row_index": 2,
        "category_row_index": 1,
        "remove_empty_cols": True,
    },
    "주차장  표준공정": {
        "skip_first_row": True,
        "header_row_index": 2,
        "category_row_index": 1,
        "remove_empty_cols": True,
    },
    "층고6.5m이상구간": {
        "skip_first_row": True,
        "header_row_index": 2,
        "category_row_index": 1,
        "remove_empty_cols": True,
    },
}


@dataclass
class SheetVerification:
    """시트별 검증 결과"""
    sheet_name: str
    excel_rows: int
    excel_cols: int
    md_rows: int
    md_cols: int
    header_loss: bool = False
    unnamed_columns: list = field(default_factory=list)
    float_precision_issues: list = field(default_factory=list)
    mismatched_cells: list = field(default_factory=list)
    accuracy_score: float = 0.0

    @property
    def row_match(self) -> bool:
        return self.excel_rows == self.md_rows

    @property
    def col_match(self) -> bool:
        return self.excel_cols == self.md_cols


def round_float_for_comparison(value) -> str:
    """부동소수점 값을 비교용 문자열로 변환 (변환 스크립트와 동일 로직)"""
    if pd.isna(value):
        return ""

    if isinstance(value, float):
        if value == int(value):
            return str(int(value))
        rounded = round(value, FLOAT_PRECISION)
        return f"{rounded:g}"

    return str(value).strip()


def get_merged_cells_map(ws) -> dict:
    """병합 셀 정보를 딕셔너리로 반환"""
    merged_map = {}
    for merged_range in ws.merged_cells.ranges:
        min_row, min_col = merged_range.min_row, merged_range.min_col
        max_row, max_col = merged_range.max_row, merged_range.max_col
        value = ws.cell(min_row, min_col).value

        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                merged_map[(row, col)] = value

    return merged_map


def remove_empty_columns(df: pd.DataFrame) -> pd.DataFrame:
    """빈 컬럼 제거 (헤더가 비어있거나 모든 데이터가 빈 컬럼)"""
    cols_to_keep = []

    for col_idx, col in enumerate(df.columns):
        col_data = df.iloc[:, col_idx]

        col_str = str(col).strip()
        if not col_str or (col_str.startswith("Col") and col_str[3:].isdigit()):
            if col_data.isna().all() or (col_data.astype(str).str.strip() == "").all():
                continue

        if col_data.isna().all():
            continue
        if (col_data.astype(str).str.strip() == "").all():
            continue

        cols_to_keep.append(col_idx)

    return df.iloc[:, cols_to_keep]


def read_excel_with_headers(wb, sheet_name: str, header_rows: int = 1) -> pd.DataFrame:
    """변환 스크립트와 동일한 방식으로 Excel 시트 읽기"""
    ws = wb[sheet_name]
    merged_map = get_merged_cells_map(ws)

    # 시트별 특수 설정 확인
    header_config = SHEET_HEADER_CONFIG.get(sheet_name, {})
    skip_first_row = header_config.get("skip_first_row", False)
    header_row_index = header_config.get("header_row_index", 0)
    category_row_index = header_config.get("category_row_index", None)
    should_remove_empty_cols = header_config.get("remove_empty_cols", False)

    data = []
    max_col = ws.max_column
    max_row = ws.max_row

    for row_idx in range(1, max_row + 1):
        row_data = []
        for col_idx in range(1, max_col + 1):
            if (row_idx, col_idx) in merged_map:
                value = merged_map[(row_idx, col_idx)]
            else:
                value = ws.cell(row_idx, col_idx).value
            row_data.append(value)
        data.append(row_data)

    if not data:
        return pd.DataFrame()

    # 특수 헤더 처리: Row 스킵 → 특정 행을 헤더로 사용
    if skip_first_row and len(data) > header_row_index:
        headers = data[header_row_index]

        # 카테고리 행 포함 여부에 따라 body_data 구성
        if category_row_index is not None and category_row_index < len(data):
            # 카테고리 행을 데이터의 첫 번째 행으로 포함
            category_row = data[category_row_index]
            body_data = [category_row] + data[header_row_index + 1:]
        else:
            body_data = data[header_row_index + 1:]

        cleaned_headers = []
        for i, h in enumerate(headers):
            if h is None or str(h).strip() == "":
                cleaned_headers.append(f"Col{i+1}")
            else:
                cleaned_headers.append(str(h).strip())

        df = pd.DataFrame(body_data, columns=cleaned_headers)

        if should_remove_empty_cols:
            df = remove_empty_columns(df)

        df = df.dropna(how='all')
        return df

    if header_rows == 0:
        df = pd.DataFrame(data)
        df.columns = [f"Col{i+1}" for i in range(len(df.columns))]
        return df

    if header_rows == 1:
        headers = data[0]
        df = pd.DataFrame(data[1:], columns=headers)
    else:
        header_data = data[:header_rows]
        body_data = data[header_rows:]

        combined_headers = []
        for col_idx in range(max_col):
            parts = []
            last_non_empty = ""
            for row_idx in range(header_rows):
                val = header_data[row_idx][col_idx]
                if val is not None and str(val).strip():
                    last_non_empty = str(val).strip()
                    parts.append(last_non_empty)

            unique_parts = []
            for p in parts:
                if p and p not in unique_parts:
                    unique_parts.append(p)

            if unique_parts:
                combined_headers.append(" / ".join(unique_parts))
            else:
                combined_headers.append(f"Col{col_idx+1}")

        df = pd.DataFrame(body_data, columns=combined_headers)

    # 컬럼명 정리
    new_columns = []
    unnamed_count = 0
    for col in df.columns:
        if col is None or str(col).strip() == "" or str(col).startswith("Unnamed:"):
            unnamed_count += 1
            new_columns.append(f"항목{unnamed_count}")
        else:
            new_columns.append(str(col).strip())
    df.columns = new_columns

    # 빈 컬럼 제거 (설정된 경우)
    if should_remove_empty_cols:
        df = remove_empty_columns(df)

    # 빈 행 제거
    df = df.dropna(how='all')

    return df


def parse_markdown_tables(md_content: str) -> dict[str, pd.DataFrame]:
    """마크다운 파일에서 테이블 추출"""
    tables = {}

    # 섹션별 분리 (## N. 시트명 패턴)
    section_pattern = r'## (\d+)\. ([^\n]+)\n'
    sections = re.split(section_pattern, md_content)

    # sections: [intro, num1, name1, content1, num2, name2, content2, ...]
    for i in range(1, len(sections) - 2, 3):
        sheet_num = sections[i]
        sheet_name = sections[i + 1].strip()
        content = sections[i + 2]

        # 테이블 추출 (| 로 시작하는 행들)
        table_lines = []
        in_table = False

        for line in content.split('\n'):
            line = line.strip()
            if line.startswith('|') and '---' not in line:
                table_lines.append(line)
                in_table = True
            elif in_table and not line.startswith('|'):
                break

        if len(table_lines) >= 1:
            # 첫 줄은 헤더
            headers = [cell.strip() for cell in table_lines[0].split('|')[1:-1]]

            # 나머지는 데이터
            data = []
            for line in table_lines[1:]:
                cells = [cell.strip() for cell in line.split('|')[1:-1]]
                if len(cells) == len(headers):
                    data.append(cells)

            if data:
                df = pd.DataFrame(data, columns=headers)
                tables[sheet_name] = df

    return tables


def is_float_precision_issue(excel_val, md_val) -> bool:
    """부동소수점 정밀도 문제 감지"""
    try:
        excel_float = float(excel_val)
        md_float = float(md_val)

        # 값이 거의 같지만 정확히 같지 않은 경우
        if excel_float != md_float and abs(excel_float - md_float) < NUMERIC_TOLERANCE:
            return True

        # 부동소수점 표현 오류 패턴 (예: 654.3000000000001)
        md_str = str(md_val)
        if re.search(r'\d+\.[\d]{10,}', md_str):
            return True

    except (ValueError, TypeError):
        pass

    return False


def compare_values(excel_val, md_val) -> tuple[bool, Optional[str]]:
    """두 값 비교 및 불일치 이유 반환"""
    # 변환 스크립트와 동일한 방식으로 정규화
    excel_normalized = round_float_for_comparison(excel_val)
    md_str = str(md_val).strip() if md_val else ""

    # 정확히 일치
    if excel_normalized == md_str:
        return True, None

    # 빈 값 비교
    if excel_normalized == "" and md_str == "":
        return True, None

    # 숫자 비교 (허용 오차 내)
    try:
        excel_float = float(excel_val) if pd.notna(excel_val) else None
        md_float = float(md_val) if md_val and md_val.strip() else None

        if excel_float is not None and md_float is not None:
            if abs(excel_float - md_float) <= NUMERIC_TOLERANCE:
                return True, None
            else:
                return False, f"숫자 불일치: {excel_float} vs {md_float}"
    except (ValueError, TypeError):
        pass

    # 문자열 불일치
    return False, f"값 불일치: '{excel_normalized}' vs '{md_str}'"


def verify_sheet(sheet_name: str, excel_df: pd.DataFrame, md_df: pd.DataFrame) -> SheetVerification:
    """단일 시트 검증"""
    result = SheetVerification(
        sheet_name=sheet_name,
        excel_rows=len(excel_df),
        excel_cols=len(excel_df.columns),
        md_rows=len(md_df),
        md_cols=len(md_df.columns)
    )

    # 헤더 손실 감지 (Unnamed: 컬럼 체크)
    md_columns = list(md_df.columns)
    unnamed_cols = [col for col in md_columns if str(col).startswith('Unnamed:')]
    if unnamed_cols:
        result.header_loss = True
        result.unnamed_columns = unnamed_cols

    # 부동소수점 정밀도 문제 감지
    for idx, row in md_df.iterrows():
        for col in md_df.columns:
            val = row[col]
            if is_float_precision_issue(val, val):
                result.float_precision_issues.append({
                    'row': idx,
                    'col': col,
                    'value': str(val)
                })

    # 셀 값 샘플링 비교
    if len(excel_df) > 0 and len(md_df) > 0:
        # 샘플 행 선택
        sample_size = min(SAMPLE_ROWS_PER_SHEET, len(excel_df), len(md_df))
        sample_indices = random.sample(range(min(len(excel_df), len(md_df))), sample_size)

        total_cells = 0
        matched_cells = 0

        for idx in sample_indices:
            excel_row = excel_df.iloc[idx]
            md_row = md_df.iloc[idx]

            for col_idx in range(min(len(excel_df.columns), len(md_df.columns))):
                excel_val = excel_row.iloc[col_idx]
                md_val = md_row.iloc[col_idx]

                total_cells += 1
                is_match, reason = compare_values(excel_val, md_val)

                if is_match:
                    matched_cells += 1
                else:
                    result.mismatched_cells.append({
                        'row': idx,
                        'col': col_idx,
                        'excel_val': str(excel_val),
                        'md_val': str(md_val),
                        'reason': reason
                    })

        if total_cells > 0:
            result.accuracy_score = (matched_cells / total_cells) * 100

    return result


def generate_report(verifications: list[SheetVerification]) -> str:
    """검증 결과 리포트 생성"""
    report_lines = [
        "# Excel ↔ Markdown 검증 리포트",
        "",
        f"- **검증 일시**: {pd.Timestamp.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"- **Excel 파일**: `{EXCEL_FILE.name}`",
        f"- **Markdown 파일**: `{MD_FILE.name}`",
        f"- **검증된 시트 수**: {len(verifications)}개",
        "",
        "---",
        "",
        "## 요약",
        "",
        "| 시트명 | 행(Excel/MD) | 열(Excel/MD) | 헤더손실 | 정밀도오류 | 정확도 |",
        "| --- | --- | --- | --- | --- | --- |",
    ]

    total_accuracy = 0
    header_loss_count = 0
    precision_issue_count = 0

    for v in verifications:
        row_status = "✅" if v.row_match else "❌"
        col_status = "✅" if v.col_match else "❌"
        header_status = "❌" if v.header_loss else "✅"

        if v.header_loss:
            header_loss_count += 1

        precision_issues = len(v.float_precision_issues)
        if precision_issues > 0:
            precision_issue_count += 1

        total_accuracy += v.accuracy_score

        report_lines.append(
            f"| {v.sheet_name} | {v.excel_rows}/{v.md_rows} {row_status} | "
            f"{v.excel_cols}/{v.md_cols} {col_status} | {header_status} | "
            f"{precision_issues}건 | {v.accuracy_score:.1f}% |"
        )

    avg_accuracy = total_accuracy / len(verifications) if verifications else 0

    report_lines.extend([
        "",
        f"**전체 평균 정확도**: {avg_accuracy:.1f}%",
        f"**헤더 손실 시트**: {header_loss_count}개",
        f"**정밀도 오류 시트**: {precision_issue_count}개",
        "",
        "---",
        "",
        "## 심각도별 문제점",
        "",
        "### 🔴 HIGH - 헤더 손실",
        "",
    ])

    # 헤더 손실 상세
    header_loss_sheets = [v for v in verifications if v.header_loss]
    if header_loss_sheets:
        for v in header_loss_sheets:
            report_lines.append(f"- **{v.sheet_name}**: `Unnamed` 컬럼 {len(v.unnamed_columns)}개")
            if v.unnamed_columns[:5]:
                report_lines.append(f"  - 예시: {', '.join(v.unnamed_columns[:5])}")
    else:
        report_lines.append("_헤더 손실 없음_")

    report_lines.extend([
        "",
        "### 🟡 MEDIUM - 부동소수점 정밀도 오류",
        "",
    ])

    # 정밀도 오류 상세
    precision_sheets = [v for v in verifications if v.float_precision_issues]
    if precision_sheets:
        for v in precision_sheets:
            report_lines.append(f"- **{v.sheet_name}**: {len(v.float_precision_issues)}건")
            for issue in v.float_precision_issues[:3]:
                report_lines.append(f"  - Row {issue['row']}, Col {issue['col']}: `{issue['value']}`")
    else:
        report_lines.append("_정밀도 오류 없음_")

    report_lines.extend([
        "",
        "### 🟢 LOW - 셀 값 불일치",
        "",
    ])

    # 셀 불일치 상세
    mismatch_sheets = [v for v in verifications if v.mismatched_cells]
    if mismatch_sheets:
        for v in mismatch_sheets:
            report_lines.append(f"- **{v.sheet_name}**: {len(v.mismatched_cells)}건")
            for mismatch in v.mismatched_cells[:3]:
                report_lines.append(
                    f"  - Row {mismatch['row']}, Col {mismatch['col']}: "
                    f"`{mismatch['excel_val'][:20]}` → `{mismatch['md_val'][:20]}`"
                )
    else:
        report_lines.append("_셀 값 불일치 없음_")

    report_lines.extend([
        "",
        "---",
        "",
        "## 권장 수정 사항",
        "",
        "1. **헤더 처리 개선**: `pd.read_excel(header=[0,1])` 또는 `skiprows` 사용",
        "2. **숫자 반올림**: 소수점 4자리로 반올림하여 정밀도 오류 방지",
        "3. **병합 셀 처리**: `openpyxl`로 병합 셀 영역 직접 파싱",
        "",
    ])

    return "\n".join(report_lines)


def main():
    print("=" * 60)
    print("🔍 Excel ↔ Markdown 데이터 정합성 검증")
    print("=" * 60)
    print()

    # Excel 파일 읽기 (openpyxl 사용)
    print(f"📂 Excel 파일: {EXCEL_FILE}")
    wb = load_workbook(EXCEL_FILE, data_only=True)
    sheet_names = wb.sheetnames
    print(f"   └─ 발견된 시트: {len(sheet_names)}개")

    # Markdown 파일 읽기
    print(f"📂 Markdown 파일: {MD_FILE}")
    md_content = MD_FILE.read_text(encoding='utf-8')
    md_tables = parse_markdown_tables(md_content)
    print(f"   └─ 추출된 테이블: {len(md_tables)}개")

    print()
    print("─" * 60)
    print()

    # 시트별 검증
    verifications = []
    random.seed(42)  # 재현 가능한 샘플링

    for sheet_name in sheet_names:
        header_rows = SHEET_HEADER_ROWS.get(sheet_name, 1)
        excel_df = read_excel_with_headers(wb, sheet_name, header_rows)

        if sheet_name in md_tables:
            md_df = md_tables[sheet_name]
            result = verify_sheet(sheet_name, excel_df, md_df)
        else:
            # MD에서 시트를 찾지 못한 경우
            result = SheetVerification(
                sheet_name=sheet_name,
                excel_rows=len(excel_df),
                excel_cols=len(excel_df.columns),
                md_rows=0,
                md_cols=0
            )

        verifications.append(result)

        # 진행 상황 출력
        status = "✅" if result.accuracy_score >= 95 else "⚠️" if result.accuracy_score >= 80 else "❌"
        print(f"{status} {sheet_name}: 정확도 {result.accuracy_score:.1f}%")

        if result.header_loss:
            print(f"   └─ ⚠️ 헤더 손실 감지: {len(result.unnamed_columns)}개 Unnamed 컬럼")
        if result.float_precision_issues:
            print(f"   └─ ⚠️ 부동소수점 오류: {len(result.float_precision_issues)}건")

    print()
    print("─" * 60)
    print()

    # 리포트 생성
    report = generate_report(verifications)
    REPORT_FILE.write_text(report, encoding='utf-8')
    print(f"📋 검증 리포트 생성: {REPORT_FILE}")

    # 전체 요약
    avg_accuracy = sum(v.accuracy_score for v in verifications) / len(verifications)
    header_loss = sum(1 for v in verifications if v.header_loss)
    precision_errors = sum(1 for v in verifications if v.float_precision_issues)

    print()
    print("=" * 60)
    print("📊 검증 결과 요약")
    print("=" * 60)
    print(f"   • 전체 평균 정확도: {avg_accuracy:.1f}%")
    print(f"   • 헤더 손실 시트: {header_loss}/{len(verifications)}개")
    print(f"   • 정밀도 오류 시트: {precision_errors}/{len(verifications)}개")

    wb.close()

    if avg_accuracy < 95 or header_loss > 0:
        print()
        print("⚠️ 변환 스크립트 개선이 필요합니다.")
    else:
        print()
        print("✅ 검증 완료: 데이터 정합성 양호")


if __name__ == "__main__":
    main()
